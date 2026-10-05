"""
Genera server/plantilla-patitas.xlsx: la hoja del servidor con todas sus pestañas,
columnas, una pestaña de Regalos lista para usar y la lista de objetos válidos.
    python3 scripts/make-sheet-template.py
Súbela a Google Drive y ábrela con Google Sheets; después añade el Apps Script.
"""
import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parent.parent
items_src = (ROOT / 'src/data/items.ts').read_text(encoding='utf8')

# --- Catálogo de objetos leído de src/data/items.ts ---
CATEGORIAS = {'food': 'Comida', 'accessory': 'Accesorio', 'furniture': 'Mueble', 'accelerator': 'Crecimiento', 'medicine': 'Salud', 'special': 'Especial'}
objetos = []
for block in re.findall(r"\{\s*id: '([^']+)',(.*?)\n  \}", items_src, re.S):
    oid, body = block
    if 'free: true' in body:
        continue
    name = re.search(r"name: '([^']+)'", body).group(1)
    cat = re.search(r"category: '([^']+)'", body).group(1)
    desc = re.search(r"description: '([^']+)'", body).group(1)
    coins = re.search(r"coins: (\d+)", body)
    stars = re.search(r"stars: (\d+)", body)
    precio = f"{coins.group(1)} monedas" if coins else f"{stars.group(1)} estrellas" if stars else ''
    objetos.append((oid, name, CATEGORIAS.get(cat, cat), precio, desc))

FONT = 'Arial'
HEAD_FILL = PatternFill('solid', fgColor='FF8A3D')
EDIT_FILL = PatternFill('solid', fgColor='FFF2CC')
AUTO_FILL = PatternFill('solid', fgColor='EEEEEE')
thin = Side(style='thin', color='E0D2C3')

def header(ws, cols, widths, auto):
    ws.append(cols)
    for i, c in enumerate(ws[1], start=1):
        c.font = Font(name=FONT, bold=True, color='FFFFFF')
        c.fill = HEAD_FILL if not auto else PatternFill('solid', fgColor='8A7566')
        c.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[c.column_letter].width = widths[i - 1]
    ws.freeze_panes = 'A2'

def text_cols(ws, letters, rows=300):
    for L in letters:
        for r in range(2, rows + 1):
            ws[f'{L}{r}'].number_format = '@'

wb = Workbook()

# ---------- Instrucciones ----------
ins = wb.active
ins.title = 'Instrucciones'
lines = [
    ('Servidor de Patitas — plantilla', True),
    ('', False),
    ('1. Sube este archivo a Google Drive y ábrelo con Google Sheets (Archivo › Guardar como Hojas de cálculo de Google).', False),
    ('2. En la hoja: Extensiones › Apps Script. Pega el contenido de server/patitas-apps-script.js y guarda.', False),
    ('3. Implementar › Nueva implementación › Aplicación web. Ejecutar como: Yo. Acceso: Cualquier usuario.', False),
    ('4. Copia la URL que termina en /exec y pégala en el juego: Ajustes › Servidor › Conectar.', False),
    ('', False),
    ('Pestañas', True),
    ('Gris = las rellena el juego automáticamente. No las edites a mano (solo para consultar).', False),
    ('Naranja = Regalos: aquí escribes tú. Celdas amarillas = las que debes rellenar.', False),
    ('', False),
    ('Cómo enviar un regalo (pestaña Regalos)', True),
    ('regalo: título que verá el jugador (ej. "Bienvenida").', False),
    ('para: el ID del jugador (Ajustes › Servidor › Tu ID) o la palabra "todos".', False),
    ('monedas / estrellas: cantidades a regalar (0 si nada).', False),
    ('objeto: id de un objeto de la pestaña Objetos (opcional). cantidad: cuántos (para comida, medicina...).', False),
    ('mensaje: texto que aparece al recibirlo.', False),
    ('activo: SI para enviarlo, NO para pausarlo.', False),
    ('entregados: NO lo toques; el servidor apunta aquí los IDs que ya lo recibieron (cada jugador lo recibe una sola vez).', False),
    ('El jugador lo recibe la próxima vez que el juego se sincroniza (al abrirlo o cada 3 minutos).', False),
]
for text, bold in lines:
    ins.append([text])
    ins.cell(ins.max_row, 1).font = Font(name=FONT, bold=bold, size=14 if bold and ins.max_row == 1 else 11)
ins.column_dimensions['A'].width = 120

# ---------- Pestañas automáticas ----------
auto_sheets = {
    'Jugadores': (['id', 'apodo', 'mascota', 'especie', 'etapa', 'sexo', 'genes', 'accesorios', 'visto', 'pendiente'], [28, 22, 16, 10, 10, 8, 40, 30, 20, 10], ['A']),
    'Adopciones': (['oferta', 'dueno', 'duenoApodo', 'mascotaId', 'nombre', 'especie', 'sexo', 'genes', 'precio', 'fecha', 'estado', 'adoptante'], [20, 28, 22, 26, 14, 10, 8, 40, 8, 20, 12, 28], ['A', 'B', 'D', 'L']),
    'Parejas': (['anuncio', 'dueno', 'duenoApodo', 'mascotaId', 'nombre', 'especie', 'sexo', 'genes', 'tarifa', 'fecha', 'usos', 'activo'], [20, 28, 22, 26, 14, 10, 8, 40, 8, 20, 6, 8], ['A', 'B', 'D']),
    'Puntajes': (['id', 'apodo', 'competencia', 'puntaje', 'estrellas', 'mascota', 'especie', 'fecha'], [28, 22, 18, 9, 9, 14, 10, 20], ['A']),
}
for name, (cols, widths, txt) in auto_sheets.items():
    ws = wb.create_sheet(name)
    header(ws, cols, widths, auto=True)
    text_cols(ws, txt)
    ws['A1'].comment = Comment('Esta pestaña la rellena el juego automáticamente. No la edites.', 'Patitas')

# ---------- Regalos (la edita el dueño) ----------
reg = wb.create_sheet('Regalos', 1)
cols = ['regalo', 'para', 'monedas', 'estrellas', 'objeto', 'cantidad', 'mensaje', 'activo', 'entregados']
header(reg, cols, [18, 28, 10, 10, 20, 10, 44, 8, 40], auto=False)
text_cols(reg, ['B', 'I'])
ejemplos = [
    ['Bienvenida', 'todos', 300, 3, 'acc_bow', 1, '¡Gracias por probar Patitas! Aquí tienes un lazo de regalo.', 'NO', ''],
    ['Regalo de prueba', 'PEGA-AQUI-EL-ID', 1000, 20, 'food_cake', 5, 'Para probar el juego sin límites.', 'NO', ''],
]
for r, row in enumerate(ejemplos, start=2):
    for c, value in enumerate(row, start=1):
        reg.cell(r, c, value)
for r in range(2, 102):
    for c in range(1, 9):
        reg.cell(r, c).fill = EDIT_FILL
        reg.cell(r, c).border = Border(bottom=thin)
    reg.cell(r, 9).fill = AUTO_FILL
for row in reg.iter_rows(min_row=1, max_row=102):
    for c in row:
        if c.row > 1:
            c.font = Font(name=FONT)
reg['B1'].comment = Comment('ID del jugador (Ajustes › Servidor › Tu ID) o "todos".', 'Patitas')
reg['E1'].comment = Comment('Opcional: id de la pestaña Objetos.', 'Patitas')
reg['H1'].comment = Comment('SI = se entrega. NO = pausado. Los ejemplos están en NO.', 'Patitas')
reg['I1'].comment = Comment('Lo rellena el servidor. No lo toques.', 'Patitas')

dv_activo = DataValidation(type='list', formula1='"SI,NO"', allow_blank=True)
reg.add_data_validation(dv_activo)
dv_activo.add('H2:H101')
dv_num = DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
reg.add_data_validation(dv_num)
dv_num.add('C2:D101')
dv_num.add('F2:F101')

# ---------- Objetos (referencia) ----------
obj = wb.create_sheet('Objetos')
header(obj, ['id', 'nombre', 'tipo', 'precio en tienda', 'para qué sirve'], [22, 24, 14, 18, 50], auto=True)
for o in objetos:
    obj.append(list(o))
for row in obj.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=FONT)
dv_obj = DataValidation(type='list', formula1=f"=Objetos!$A$2:$A${len(objetos) + 1}", allow_blank=True)
reg.add_data_validation(dv_obj)
dv_obj.add('E2:E101')

for ws in wb.worksheets:
    for row in ws.iter_rows(min_row=1, max_row=1):
        for c in row:
            if c.font.name != FONT:
                c.font = Font(name=FONT, bold=True, color=c.font.color)

out = ROOT / 'server' / 'plantilla-patitas.xlsx'
wb.save(out)
print(out, f'({len(objetos)} objetos)')
